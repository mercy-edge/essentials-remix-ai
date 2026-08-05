#!/usr/bin/env python3
"""
Generate the project-wide Run & Bun-style documentation workbook.

Available Pokémon is the canonical sprite / type / BST catalog.
Wild Encounters, Trainer Battles, and Static Battles reference that catalog
(local sprites in .xlsx / HTML; INDEX/MATCH formulas in Google Sheets).

Outputs:
  docs/sheets/Essentials_Remix_AI_Project.xlsx
  docs/sheets/Essentials_Remix_AI_Pokemon.xlsx  (alias of the same workbook)
  docs/sheets/Available_Pokemon.html

Usage:
  python3 tools/generate_project_docs.py
  python3 tools/generate_project_docs.py --no-images
"""

from __future__ import annotations

import argparse
import base64
import io
import shutil
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import docs_lib as d

ROOT = d.ROOT


def _embed_icon(ws, row: int, col_letter: str, stem: str, size: int = d.ICON_SIZE):
    png = d.icon_png_bytes(stem, size)
    if not png:
        return
    img = XLImage(io.BytesIO(png))
    img.width = size
    img.height = size
    img.anchor = f"{col_letter}{row}"
    ws.add_image(img)


def write_index_sheet(
    wb: Workbook,
    *,
    base_count: int,
    form_count: int,
    poke_total: int,
    encounter_count: int,
    trainer_count: int,
    static_count: int,
):
    ws = wb.active
    ws.title = "Index"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    ws["A1"] = d.PROJECT_TITLE
    ws["A1"].font = d.FONT_TITLE
    ws["A1"].fill = d.FILL_TITLE
    ws["A1"].alignment = d.LEFT
    ws.row_dimensions[1].height = 32

    lines = [
        ("Generated", str(date.today())),
        ("Scope", "Whole-project docs hub (Pokémon catalog → encounters → trainers)"),
        ("", ""),
        ("Available Pokémon", f"{poke_total} rows ({base_count} base + {form_count} forms)"),
        ("Wild Encounters", f"{encounter_count} slots — sprites/types/BST via catalog"),
        ("Trainer Battles", f"{trainer_count} trainers — lead sprite + party from catalog"),
        ("Trainer Party", "One row per party member with catalog sprite lookup"),
        ("Static Battles", f"{static_count} scripted wild battles"),
        ("", ""),
        (
            "Sprite sheet",
            f"'{d.POKEMON_TAB}' is canonical. Other tabs look up Species ID → Sprite "
            f"(xlsx embeds icons; Google Sheets uses INDEX/MATCH into that tab).",
        ),
        (
            "Regenerate",
            "python3 tools/generate_project_docs.py\n"
            "python3 tools/sync_project_google_sheets.py",
        ),
        (
            "Offline HTML",
            "Available_Pokemon.html — filterable catalog with local sprites",
        ),
    ]
    r = 3
    for label, value in lines:
        ws.cell(r, 1, label).font = Font(name="Arial", size=10, bold=True)
        cell = ws.cell(r, 2, value)
        cell.font = d.FONT_BODY
        cell.alignment = d.WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        r += 1
    ws.column_dimensions["A"].width = 18
    for col in range(2, 7):
        ws.column_dimensions[get_column_letter(col)].width = 18


def write_pokemon_sheet(wb: Workbook, rows: list[d.SpeciesRow], embed_images: bool):
    ws = wb.create_sheet(d.POKEMON_TAB, 1)
    headers = [
        "Sprite",
        "#",
        "Name",
        "Form",
        "Type 1",
        "Type 2",
        "HP",
        "Atk",
        "Def",
        "SpA",
        "SpD",
        "Spe",
        "BST",
        "Generation",
        "Abilities",
        "Hidden Ability",
        "Flags",
        "Species ID",
    ]
    widths = [8, 6, 18, 22, 10, 10, 6, 6, 6, 6, 6, 6, 7, 10, 28, 16, 14, 16]
    d.apply_header_row(ws, 1, headers, widths)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:R{len(rows) + 1}"

    for i, row in enumerate(rows):
        r = i + 2
        alt = i % 2 == 1
        fill = d.FILL_ALT if alt else d.FILL_WHITE
        values = [
            "",
            row.dex,
            row.name,
            row.form_name,
            row.types[0].title() if row.types else "",
            row.types[1].title() if len(row.types) > 1 else "",
            *row.stats,
            row.bst,
            row.generation,
            row.abilities,
            row.hidden_ability.replace(",", " / ") if row.hidden_ability else "",
            row.flags,
            row.catalog_key,
        ]
        ws.row_dimensions[r].height = 32 if embed_images else 18
        for col, val in enumerate(values, 1):
            cell = ws.cell(r, col, val)
            cell.font = d.FONT_BODY
            cell.fill = fill
            cell.border = d.BORDER_ALL
            cell.alignment = d.CENTER if col not in (3, 4, 15, 16) else d.LEFT
        for tcol, tidx in ((5, 0), (6, 1)):
            if tidx < len(row.types):
                c = ws.cell(r, tcol)
                c.fill = d.type_fill(row.types[tidx])
                c.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        ws.cell(r, 13).font = Font(name="Arial", size=10, bold=True)
        if embed_images:
            _embed_icon(ws, r, "A", row.sprite_key)


def write_by_type_sheet(wb: Workbook, rows: list[d.SpeciesRow]):
    ws = wb.create_sheet("By Type")
    headers = ["Type", "Count", "Avg BST", "Examples (highest BST)"]
    d.apply_header_row(ws, 1, headers, [12, 10, 10, 60])
    buckets: dict[str, list[d.SpeciesRow]] = {t: [] for t in d.TYPE_COLORS}
    for row in rows:
        for t in row.types:
            buckets.setdefault(t, []).append(row)
    r = 2
    for t in d.TYPE_COLORS:
        group = buckets.get(t, [])
        if not group:
            continue
        avg = round(sum(x.bst for x in group) / len(group), 1)
        top = sorted(group, key=lambda x: (-x.bst, x.name))[:8]
        examples = ", ".join(f"{x.display_name} ({x.bst})" for x in top)
        for col, val in enumerate([t.title(), len(group), avg, examples], 1):
            cell = ws.cell(r, col, val)
            cell.border = d.BORDER_ALL
            cell.font = d.FONT_BODY
            if col == 1:
                cell.fill = d.type_fill(t)
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.alignment = d.CENTER
        r += 1


def write_bst_tiers_sheet(wb: Workbook, rows: list[d.SpeciesRow]):
    ws = wb.create_sheet("BST Tiers")
    tiers = [
        ("NU / weak", 0, 399),
        ("RU / mid", 400, 479),
        ("UU / strong", 480, 529),
        ("OU / very strong", 530, 569),
        ("Uber / legendary tier", 570, 9999),
    ]
    d.apply_header_row(ws, 1, ["Tier", "BST range", "Count", "Examples"], [22, 12, 10, 70])
    for i, (label, lo, hi) in enumerate(tiers):
        group = [x for x in rows if lo <= x.bst <= hi]
        examples = ", ".join(
            x.display_name for x in sorted(group, key=lambda x: (-x.bst, x.name))[:12]
        )
        for col, val in enumerate(
            [label, f"{lo}–{hi if hi < 9999 else '∞'}", len(group), examples], 1
        ):
            cell = ws.cell(i + 2, col, val)
            cell.border = d.BORDER_ALL
            cell.font = d.FONT_BODY


def write_encounters_sheet(
    wb: Workbook,
    slots: list[d.EncounterSlot],
    catalog: dict[str, d.SpeciesRow],
    maps: dict,
    embed_images: bool,
    sheets_formulas: bool,
):
    ws = wb.create_sheet("Wild Encounters")
    headers = [
        "Location",
        "Map",
        "Method",
        "Sprite",
        "Pokémon",
        "Type",
        "BST",
        "Level",
        "Rate %",
        "Species ID",
        "Notes",
    ]
    widths = [28, 8, 16, 8, 18, 14, 7, 10, 10, 14, 20]
    d.apply_header_row(ws, 1, headers, widths)
    ws.freeze_panes = "A2"

    groups: dict[tuple, list[d.EncounterSlot]] = defaultdict(list)
    for s in slots:
        groups[(s.map_id, s.version or 0, s.location)].append(s)

    row = 2
    for key in sorted(groups.keys(), key=lambda k: (d.map_sort_key(k[0], maps), k[1], k[2])):
        map_id, _ver, location = key
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        sec = ws.cell(row, 1, f"▸  {location}  (Map {map_id})")
        sec.font = d.FONT_SECTION
        sec.fill = d.FILL_SECTION
        ws.row_dimensions[row].height = 22
        row += 1

        by_method: dict[str, list[d.EncounterSlot]] = defaultdict(list)
        for s in groups[key]:
            by_method[s.method].append(s)
        alt = False
        for method in sorted(
            by_method.keys(),
            key=lambda m: (d.METHOD_ORDER.index(m) if m in d.METHOD_ORDER else 99, m),
        ):
            for s in sorted(by_method[method], key=lambda x: (-x.rate, x.species)):
                sp = d.lookup_species(catalog, s.species)
                types = sp.types if sp else []
                catalog_key = sp.catalog_key if sp else s.species
                sprite_val = d.sheets_sprite_formula(catalog_key) if sheets_formulas else ""
                values = [
                    location,
                    map_id,
                    d.METHOD_LABELS.get(method, method),
                    sprite_val,
                    sp.display_name if sp else d.fmt_name(s.species),
                    d.fmt_types(types) if types else "?",
                    sp.bst if sp else "",
                    d.level_str(s.level_min, s.level_max),
                    s.rate,
                    catalog_key,
                    "",
                ]
                ws.row_dimensions[row].height = 32 if embed_images and not sheets_formulas else 18
                for col, val in enumerate(values, 1):
                    cell = ws.cell(row, col, val)
                    d.style_data_cell(cell, alt=alt)
                    if col == 6 and types:
                        cell.fill = d.type_fill(types[0])
                        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                    if col == 7 and sp:
                        cell.font = Font(name="Arial", size=10, bold=True)
                if embed_images and not sheets_formulas and sp:
                    _embed_icon(ws, row, "D", sp.sprite_key)
                alt = not alt
                row += 1


def write_encounter_summary(wb: Workbook, slots: list[d.EncounterSlot], catalog, maps):
    ws = wb.create_sheet("Encounter Summary")
    headers = ["Location", "Map", "Methods", "Species (unique)", "Level Range", "Slot Count"]
    d.apply_header_row(ws, 1, headers, [28, 8, 24, 40, 14, 12])
    ws.freeze_panes = "A2"
    by_loc: dict[tuple, list[d.EncounterSlot]] = defaultdict(list)
    for s in slots:
        by_loc[(s.map_id, s.location)].append(s)
    row = 2
    alt = False
    for map_id, location in sorted(by_loc.keys(), key=lambda k: d.map_sort_key(k[0], maps)):
        group = by_loc[(map_id, location)]
        methods = sorted({d.METHOD_LABELS.get(s.method, s.method) for s in group})
        species = sorted(
            {
                (d.lookup_species(catalog, s.species).display_name
                 if d.lookup_species(catalog, s.species)
                 else d.fmt_name(s.species))
                for s in group
            }
        )
        levels = [s.level_min for s in group] + [s.level_max for s in group if s.level_max]
        lv_range = f"{min(levels)}–{max(levels)}" if levels else "—"
        for col, val in enumerate(
            [location, map_id, ", ".join(methods), ", ".join(species), lv_range, len(group)],
            1,
        ):
            d.style_data_cell(ws.cell(row, col, val), alt=alt)
        alt = not alt
        row += 1


def write_trainers_sheet(
    wb: Workbook,
    trainers: list[d.Trainer],
    trainer_types: dict[str, str],
    catalog: dict[str, d.SpeciesRow],
    placements: dict,
    maps: dict,
    embed_images: bool,
    sheets_formulas: bool,
):
    ws = wb.create_sheet("Trainer Battles")
    headers = [
        "Location",
        "Map",
        "Trainer",
        "Lead Sprite",
        "Party",
        "Trainer Items",
        "Lose Text",
        "Lead Species ID",
        "Notes",
    ]
    widths = [26, 8, 24, 8, 56, 22, 28, 14, 16]
    d.apply_header_row(ws, 1, headers, widths)
    ws.freeze_panes = "A2"

    placed: dict[tuple, list[dict]] = defaultdict(list)
    for t in trainers:
        key = d.trainer_key(t)
        if key in placements:
            placed[key].extend(placements[key])

    entries: list[tuple] = []
    seen: set[tuple] = set()
    for t in trainers:
        key = d.trainer_key(t)
        if key in seen:
            continue
        seen.add(key)
        locs = placed.get(key, [])
        if locs:
            for loc in locs:
                entries.append((loc["map_id"], t, loc))
        else:
            entries.append((9999, t, None))
    entries.sort(
        key=lambda e: (
            d.map_sort_key(e[0], maps) if e[0] != 9999 else (9, 99, 99, 9999, ""),
            e[1].trainer_type,
            e[1].name,
            e[1].version,
        )
    )

    row = 2
    current_map: int | None = None
    alt = False
    for map_id, trainer, loc in entries:
        if map_id != current_map:
            current_map = map_id
            loc_name = (
                maps.get(map_id, {}).get("Name", maps.get(map_id, {}).get("editor_name", f"Map {map_id}"))
                if map_id != 9999
                else "Unplaced / Script-Only Battles"
            )
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
            sec = ws.cell(row, 1, f"▸  {loc_name}")
            sec.font = d.FONT_SECTION
            sec.fill = d.FILL_SECTION
            ws.row_dimensions[row].height = 22
            row += 1
            alt = False

        is_boss = trainer.trainer_type.startswith("LEADER_") or trainer.trainer_type == "CHAMPION"
        meta = maps.get(map_id, {}) if map_id != 9999 else {}
        loc_name = meta.get("Name", "") if map_id != 9999 else "—"
        party = "\n".join(d.mon_detail_line(m, catalog) for m in trainer.pokemon)
        lead = d.lookup_species(catalog, trainer.pokemon[0].species) if trainer.pokemon else None
        lead_key = lead.catalog_key if lead else ""
        sprite_val = d.sheets_sprite_formula(lead_key) if sheets_formulas and lead_key else ""
        notes = "Boss" if is_boss else ""
        if loc and loc.get("event_id"):
            notes = (notes + f" Event {loc['event_id']}").strip()
        values = [
            loc_name,
            map_id if map_id != 9999 else "—",
            d.trainer_display_name(trainer, trainer_types),
            sprite_val,
            party,
            ", ".join(d.fmt_name(i) for i in trainer.items) if trainer.items else "—",
            trainer.lose_text or "—",
            lead_key,
            notes,
        ]
        ws.row_dimensions[row].height = 48 if embed_images and not sheets_formulas else 36
        for col, val in enumerate(values, 1):
            cell = ws.cell(row, col, val)
            fill = d.FILL_BOSS if is_boss else (d.FILL_ALT if alt else d.FILL_WHITE)
            cell.fill = fill
            cell.font = d.FONT_TRAINER if col == 3 else d.FONT_BODY
            cell.border = d.BORDER_ALL
            cell.alignment = d.WRAP
        if embed_images and not sheets_formulas and lead:
            _embed_icon(ws, row, "D", lead.sprite_key)
        alt = not alt
        row += 1


def write_trainer_party_sheet(
    wb: Workbook,
    trainers: list[d.Trainer],
    trainer_types: dict[str, str],
    catalog: dict[str, d.SpeciesRow],
    embed_images: bool,
    sheets_formulas: bool,
):
    """Expanded party rows — ready for future deep-links into the sprite catalog."""
    ws = wb.create_sheet("Trainer Party")
    headers = [
        "Trainer",
        "Slot",
        "Sprite",
        "Pokémon",
        "Type",
        "BST",
        "Level",
        "Item",
        "Moves",
        "Species ID",
    ]
    widths = [28, 6, 8, 20, 14, 7, 8, 14, 40, 14]
    d.apply_header_row(ws, 1, headers, widths)
    ws.freeze_panes = "A2"
    row = 2
    alt = False
    for trainer in trainers:
        tname = d.trainer_display_name(trainer, trainer_types)
        for slot, mon in enumerate(trainer.pokemon, 1):
            sp = d.lookup_species(catalog, mon.species)
            key = sp.catalog_key if sp else mon.species
            sprite_val = d.sheets_sprite_formula(key) if sheets_formulas else ""
            values = [
                tname,
                slot,
                sprite_val,
                (mon.name or (sp.display_name if sp else d.fmt_name(mon.species))),
                d.fmt_types(sp.types) if sp else "?",
                sp.bst if sp else "",
                mon.level,
                d.fmt_item(mon.item) or "—",
                d.fmt_moves(mon.moves),
                key,
            ]
            ws.row_dimensions[row].height = 32 if embed_images and not sheets_formulas else 18
            for col, val in enumerate(values, 1):
                cell = ws.cell(row, col, val)
                d.style_data_cell(cell, alt=alt)
                if col == 5 and sp:
                    cell.fill = d.type_fill(sp.types[0])
                    cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            if embed_images and not sheets_formulas and sp:
                _embed_icon(ws, row, "C", sp.sprite_key)
            alt = not alt
            row += 1


def write_static_sheet(
    wb: Workbook,
    static: list,
    catalog: dict[str, d.SpeciesRow],
    maps: dict,
    embed_images: bool,
    sheets_formulas: bool,
):
    ws = wb.create_sheet("Static Battles")
    headers = ["Location", "Map", "Sprite", "Pokémon", "Type", "BST", "Level", "Species ID", "Notes"]
    widths = [28, 8, 8, 18, 14, 7, 10, 14, 30]
    d.apply_header_row(ws, 1, headers, widths)
    ws.freeze_panes = "A2"
    row = 2
    alt = False
    for b in sorted(static, key=lambda x: x["map_id"]):
        map_id = b["map_id"]
        meta = maps.get(map_id, {})
        loc = meta.get("Name", meta.get("editor_name", f"Map {map_id}"))
        sp = d.lookup_species(catalog, b["species"])
        key = sp.catalog_key if sp else b["species"]
        sprite_val = d.sheets_sprite_formula(key) if sheets_formulas else ""
        types = sp.types if sp else []
        values = [
            loc,
            map_id,
            sprite_val,
            sp.display_name if sp else d.fmt_name(b["species"]),
            d.fmt_types(types) if types else "?",
            sp.bst if sp else "",
            b["level"],
            key,
            f"Event {b['event_id']} — WildBattle.start",
        ]
        ws.row_dimensions[row].height = 32 if embed_images and not sheets_formulas else 18
        for col, val in enumerate(values, 1):
            cell = ws.cell(row, col, val)
            d.style_data_cell(cell, alt=alt)
            if col == 5 and types:
                cell.fill = d.type_fill(types[0])
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        if embed_images and not sheets_formulas and sp:
            _embed_icon(ws, row, "C", sp.sprite_key)
        alt = not alt
        row += 1


def generate_html(rows: list[d.SpeciesRow]) -> str:
    cards = []
    for row in rows:
        png = d.icon_png_bytes(row.sprite_key, 64)
        if png:
            src = f"data:image/png;base64,{base64.b64encode(png).decode('ascii')}"
            img = f'<img class="sprite" src="{src}" alt="{row.name}" width="64" height="64">'
        else:
            img = '<div class="sprite missing">?</div>'
        type_html = "".join(
            f'<span class="type" style="background:#{d.TYPE_COLORS.get(t, "9CA3AF")}">{t.title()}</span>'
            for t in row.types
        )
        search = " ".join(
            [row.name, row.form_name, *row.types, row.species_id, row.flags, str(row.bst)]
        ).lower()

        def esc(s: str) -> str:
            return (
                s.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
                .replace('"', "&quot;")
            )

        cards.append(
            f"""<tr data-search="{esc(search)}">
  <td class="c">{img}</td>
  <td class="c">{row.dex}</td>
  <td><div class="name">{esc(row.name)}</div><div class="form">{esc(row.form_name)}</div></td>
  <td>{type_html}</td>
  <td class="c">{row.stats[0]}</td><td class="c">{row.stats[1]}</td><td class="c">{row.stats[2]}</td>
  <td class="c">{row.stats[3]}</td><td class="c">{row.stats[4]}</td><td class="c">{row.stats[5]}</td>
  <td class="c bst">{row.bst}</td>
</tr>"""
        )
    body = "\n".join(cards)
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{d.PROJECT_TITLE} — Available Pokémon</title>
<style>
  :root {{ --bg:#0f172a; --panel:#1e293b; --text:#e2e8f0; --muted:#94a3b8; --line:#334155; --accent:#38bdf8; }}
  * {{ box-sizing:border-box; }}
  body {{ margin:0; font-family:"Segoe UI",Tahoma,sans-serif; background:linear-gradient(180deg,#0b1224,#111827 40%,#0f172a); color:var(--text); min-height:100vh; }}
  header {{ padding:28px 24px 16px; border-bottom:1px solid var(--line); background:rgba(15,23,42,.9); position:sticky; top:0; z-index:5; backdrop-filter:blur(8px); }}
  h1 {{ margin:0 0 6px; font-size:1.6rem; }}
  .sub {{ color:var(--muted); font-size:.95rem; margin-bottom:14px; }}
  .controls {{ display:flex; flex-wrap:wrap; gap:10px; align-items:center; }}
  input[type="search"] {{ flex:1; min-width:220px; padding:10px 12px; border-radius:8px; border:1px solid var(--line); background:var(--panel); color:var(--text); }}
  .count {{ color:var(--muted); font-size:.9rem; }}
  main {{ padding:16px 24px 48px; overflow-x:auto; }}
  table {{ width:100%; border-collapse:collapse; min-width:860px; }}
  th {{ position:sticky; top:118px; background:#1e293b; text-align:left; padding:10px 8px; border-bottom:2px solid var(--accent); font-size:.85rem; }}
  td {{ padding:8px; border-bottom:1px solid var(--line); vertical-align:middle; }}
  tr:hover td {{ background:rgba(56,189,248,.08); }}
  .c {{ text-align:center; }} .name {{ font-weight:700; }} .form {{ color:var(--muted); font-size:.8rem; }}
  .bst {{ font-weight:800; color:#fbbf24; }}
  .type {{ display:inline-block; min-width:64px; text-align:center; margin:0 3px 3px 0; padding:3px 8px; border-radius:4px; font-size:.75rem; font-weight:700; color:#fff; text-shadow:0 1px 1px rgba(0,0,0,.35); }}
  .sprite {{ image-rendering:pixelated; background:#0b1224; border-radius:8px; }}
  .sprite.missing {{ width:64px; height:64px; display:grid; place-items:center; background:#334155; border-radius:8px; color:var(--muted); }}
  footer {{ padding:12px 24px 28px; color:var(--muted); font-size:.85rem; }}
</style>
</head>
<body>
<header>
  <h1>{d.PROJECT_TITLE}</h1>
  <div class="sub">Canonical Pokémon catalog (sprites · types · BST) — referenced by encounters &amp; trainers · {date.today()}</div>
  <div class="controls">
    <input id="q" type="search" placeholder="Filter by name, type, BST, flags…">
    <span class="count"><span id="shown">{len(rows)}</span> / {len(rows)} shown</span>
  </div>
</header>
<main>
<table>
  <thead><tr>
    <th>Sprite</th><th>#</th><th>Name</th><th>Types</th>
    <th>HP</th><th>Atk</th><th>Def</th><th>SpA</th><th>SpD</th><th>Spe</th><th>BST</th>
  </tr></thead>
  <tbody id="rows">
{body}
  </tbody>
</table>
</main>
<footer>
  Project docs hub: Available Pokémon is the sprite sheet other tabs reference.
  Full workbook: Essentials_Remix_AI_Project.xlsx
</footer>
<script>
const q=document.getElementById('q'), shown=document.getElementById('shown'), rows=[...document.querySelectorAll('#rows tr')];
q.addEventListener('input',()=>{{const term=q.value.trim().toLowerCase(); let n=0; for(const tr of rows){{const ok=!term||tr.dataset.search.includes(term); tr.hidden=!ok; if(ok)n++;}} shown.textContent=n;}});
</script>
</body>
</html>
"""


def build_workbook(
    *,
    poke_rows: list[d.SpeciesRow],
    base_count: int,
    form_count: int,
    catalog: dict[str, d.SpeciesRow],
    encounters: list[d.EncounterSlot],
    trainers: list[d.Trainer],
    trainer_types: dict[str, str],
    placements: dict,
    static: list,
    maps: dict,
    embed_images: bool,
    sheets_formulas: bool,
) -> Workbook:
    wb = Workbook()
    write_index_sheet(
        wb,
        base_count=base_count,
        form_count=form_count,
        poke_total=len(poke_rows),
        encounter_count=len(encounters),
        trainer_count=len(trainers),
        static_count=len(static),
    )
    write_pokemon_sheet(wb, poke_rows, embed_images=embed_images and not sheets_formulas)
    # For Google Sheets export path, Available Pokémon still needs IMAGE() formulas
    if sheets_formulas:
        ws = wb[d.POKEMON_TAB]
        for i, row in enumerate(poke_rows):
            ws.cell(i + 2, 1, f'=IMAGE("{d.pokeapi_sprite_url(row.dex)}")')
    write_by_type_sheet(wb, poke_rows)
    write_bst_tiers_sheet(wb, poke_rows)
    write_encounters_sheet(
        wb, encounters, catalog, maps, embed_images=embed_images, sheets_formulas=sheets_formulas
    )
    write_encounter_summary(wb, encounters, catalog, maps)
    write_trainers_sheet(
        wb,
        trainers,
        trainer_types,
        catalog,
        placements,
        maps,
        embed_images=embed_images,
        sheets_formulas=sheets_formulas,
    )
    write_trainer_party_sheet(
        wb,
        trainers,
        trainer_types,
        catalog,
        embed_images=embed_images,
        sheets_formulas=sheets_formulas,
    )
    write_static_sheet(
        wb, static, catalog, maps, embed_images=embed_images, sheets_formulas=sheets_formulas
    )
    return wb


def load_all():
    base, forms, poke_rows = d.collect_species_rows()
    catalog = d.build_catalog(poke_rows)
    maps = d.parse_map_metadata(d.PBS / "map_metadata.txt")
    encounters = d.parse_encounters(d.PBS / "encounters.txt", maps)
    trainers = d.parse_trainers(d.PBS / "trainers.txt")
    trainer_types = d.parse_trainer_types(d.PBS / "trainer_types.txt")
    placements, static = d.parse_battle_placements(d.PBS / "MapEvents")
    return {
        "base": base,
        "forms": forms,
        "poke_rows": poke_rows,
        "catalog": catalog,
        "maps": maps,
        "encounters": encounters,
        "trainers": trainers,
        "trainer_types": trainer_types,
        "placements": placements,
        "static": static,
    }


def generate(*, embed_images: bool = True, html: bool = True, xlsx: bool = True) -> dict:
    data = load_all()
    d.OUT_DIR.mkdir(parents=True, exist_ok=True)

    if xlsx:
        wb = build_workbook(
            poke_rows=data["poke_rows"],
            base_count=len(data["base"]),
            form_count=len(data["forms"]),
            catalog=data["catalog"],
            encounters=data["encounters"],
            trainers=data["trainers"],
            trainer_types=data["trainer_types"],
            placements=data["placements"],
            static=data["static"],
            maps=data["maps"],
            embed_images=embed_images,
            sheets_formulas=False,
        )
        wb.save(d.OUT_XLSX)
        shutil.copyfile(d.OUT_XLSX, d.OUT_XLSX_POKEMON_ALIAS)
        print(f"Wrote {d.OUT_XLSX.relative_to(ROOT)}")
        print(f"Wrote {d.OUT_XLSX_POKEMON_ALIAS.relative_to(ROOT)} (alias)")

    if html:
        d.OUT_HTML.write_text(generate_html(data["poke_rows"]), encoding="utf-8")
        print(f"Wrote {d.OUT_HTML.relative_to(ROOT)}")

    print(
        f"Done — Pokémon {len(data['poke_rows'])} | "
        f"encounters {len(data['encounters'])} | "
        f"trainers {len(data['trainers'])} | "
        f"static {len(data['static'])}"
    )
    return data


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--no-images", action="store_true")
    ap.add_argument("--xlsx-only", action="store_true")
    ap.add_argument("--html-only", action="store_true")
    args = ap.parse_args()
    generate(
        embed_images=not args.no_images,
        html=not args.xlsx_only,
        xlsx=not args.html_only,
    )


if __name__ == "__main__":
    main()
