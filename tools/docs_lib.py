#!/usr/bin/env python3
"""
Shared library for Essentials Remix AI project documentation.

The Available Pokémon catalog is the canonical sprite / type / BST source.
Encounter and trainer tabs look up species through this catalog so the whole
project docs set stays consistent as content expands.
"""

from __future__ import annotations

import io
import json
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parents[1]
PBS = ROOT / "PBS"
GRAPHICS = ROOT / "Graphics" / "Pokemon"
OUT_DIR = ROOT / "docs" / "sheets"

# Canonical project outputs
OUT_XLSX = OUT_DIR / "Essentials_Remix_AI_Project.xlsx"
OUT_HTML = OUT_DIR / "Available_Pokemon.html"
# Back-compat alias copied alongside the project workbook
OUT_XLSX_POKEMON_ALIAS = OUT_DIR / "Essentials_Remix_AI_Pokemon.xlsx"

PROJECT_TITLE = "Essentials Remix AI — Project Docs"
POKEMON_TAB = "Available Pokémon"
# Column letters on Available Pokémon (1-indexed): Sprite=A … Species ID=R
POKEMON_SPRITE_COL = "A"
POKEMON_SPECIES_ID_COL = "R"

TYPE_COLORS = {
    "NORMAL": "A8A878",
    "FIRE": "F08030",
    "WATER": "6890F0",
    "ELECTRIC": "F8D030",
    "GRASS": "78C850",
    "ICE": "98D8D8",
    "FIGHTING": "C03028",
    "POISON": "A040A0",
    "GROUND": "E0C068",
    "FLYING": "A890F0",
    "PSYCHIC": "F85888",
    "BUG": "A8B820",
    "ROCK": "B8A038",
    "GHOST": "705898",
    "DRAGON": "7038F8",
    "DARK": "705848",
    "STEEL": "B8B8D0",
    "FAIRY": "EE99AC",
}

METHOD_LABELS = {
    "Land": "Grass",
    "LandNight": "Grass (Night)",
    "LandMorning": "Grass (Morning)",
    "LandDay": "Grass (Day)",
    "Water": "Surf",
    "OldRod": "Old Rod",
    "GoodRod": "Good Rod",
    "SuperRod": "Super Rod",
    "Cave": "Cave",
    "RockSmash": "Rock Smash",
    "HeadbuttLow": "Headbutt (Low)",
    "HeadbuttHigh": "Headbutt (High)",
    "PokeRadar": "Poké Radar",
    "BugContest": "Bug Contest",
}
METHOD_ORDER = list(METHOD_LABELS.keys()) + ["Other"]

FONT_TITLE = Font(name="Arial", size=18, bold=True, color="FFFFFF")
FONT_SECTION = Font(name="Arial", size=11, bold=True, color="FFFFFF")
FONT_HEADER = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FONT_BODY = Font(name="Arial", size=10)
FONT_TRAINER = Font(name="Arial", size=11, bold=True, color="1F2937")
FONT_NOTE = Font(name="Arial", size=9, italic=True, color="6B7280")

FILL_TITLE = PatternFill("solid", fgColor="1E3A5F")
FILL_SECTION = PatternFill("solid", fgColor="2563EB")
FILL_HEADER = PatternFill("solid", fgColor="374151")
FILL_ALT = PatternFill("solid", fgColor="F3F4F6")
FILL_WHITE = PatternFill("solid", fgColor="FFFFFF")
FILL_BOSS = PatternFill("solid", fgColor="FEF3C7")

THIN = Side(style="thin", color="D1D5DB")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

ICON_SIZE = 40


@dataclass
class SpeciesRow:
    dex: int
    species_id: str
    form: int
    name: str
    form_name: str
    types: list[str]
    stats: list[int]
    bst: int
    generation: str
    flags: str
    abilities: str
    hidden_ability: str
    sprite_key: str

    @property
    def catalog_key(self) -> str:
        """Stable key used by other tabs to look up this row."""
        return self.species_id if self.form == 0 else f"{self.species_id},{self.form}"

    @property
    def display_name(self) -> str:
        if self.form_name:
            return f"{self.name} ({self.form_name})"
        return self.name


@dataclass
class EncounterSlot:
    map_id: int
    version: int | None
    location: str
    method: str
    rate: int
    species: str
    level_min: int
    level_max: int | None


@dataclass
class TrainerMon:
    species: str
    level: int
    name: str | None = None
    moves: list[str] = field(default_factory=list)
    item: str | None = None
    ability_index: str | None = None
    gender: str | None = None
    ivs: str | None = None
    shiny: bool = False
    shadow: bool = False
    ball: str | None = None


@dataclass
class Trainer:
    trainer_type: str
    name: str
    version: int
    lose_text: str = ""
    items: list[str] = field(default_factory=list)
    pokemon: list[TrainerMon] = field(default_factory=list)


def parse_pbs_sections(path: Path) -> list[tuple[str, dict[str, str]]]:
    if not path.is_file():
        return []
    sections: list[tuple[str, dict[str, str]]] = []
    current_id: str | None = None
    current: dict[str, str] = {}
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        m = re.match(r"^\[(.+)\]$", line)
        if m:
            if current_id is not None:
                sections.append((current_id, current))
            current_id = m.group(1).strip()
            current = {}
            continue
        if current_id is None or "=" not in line:
            continue
        key, val = [p.strip() for p in line.split("=", 1)]
        current[key] = val
    if current_id is not None:
        sections.append((current_id, current))
    return sections


def parse_stats(raw: str | None) -> list[int] | None:
    if not raw:
        return None
    parts = [p.strip() for p in raw.split(",")]
    if len(parts) != 6:
        return None
    try:
        return [int(p) for p in parts]
    except ValueError:
        return None


def parse_types(raw: str | None) -> list[str]:
    if not raw:
        return []
    return [t.strip().upper() for t in raw.split(",") if t.strip()]


def sprite_stem(species_id: str, form: int) -> str:
    return species_id if form == 0 else f"{species_id}_{form}"


def resolve_sprite_path(stem: str, kind: str = "Icons") -> Path | None:
    folder = GRAPHICS / kind
    for c in (folder / f"{stem}.png", folder / f"{stem.split('_')[0]}.png"):
        if c.is_file():
            return c
    return None


def load_base_species() -> list[SpeciesRow]:
    rows: list[SpeciesRow] = []
    dex = 0
    for path in (PBS / "pokemon.txt", PBS / "pokemon_base_Gen_9_Pack.txt"):
        for sid, fields in parse_pbs_sections(path):
            if "," in sid:
                continue
            stats = parse_stats(fields.get("BaseStats"))
            if not stats:
                continue
            dex += 1
            rows.append(
                SpeciesRow(
                    dex=dex,
                    species_id=sid,
                    form=0,
                    name=fields.get("Name") or sid.title(),
                    form_name="",
                    types=parse_types(fields.get("Types")),
                    stats=stats,
                    bst=sum(stats),
                    generation=fields.get("Generation", ""),
                    flags=fields.get("Flags", ""),
                    abilities=fields.get("Abilities", "").replace(",", " / "),
                    hidden_ability=fields.get("HiddenAbilities", ""),
                    sprite_key=sid,
                )
            )
    return rows


def load_forms(base_by_id: dict[str, SpeciesRow]) -> list[SpeciesRow]:
    forms: list[SpeciesRow] = []
    for path in (PBS / "pokemon_forms.txt", PBS / "pokemon_forms_Gen_9_Pack.txt"):
        for section_id, fields in parse_pbs_sections(path):
            m = re.match(r"^([A-Z0-9_]+),(\d+)$", section_id)
            if not m:
                continue
            sid, form_n = m.group(1), int(m.group(2))
            base = base_by_id.get(sid)
            if not base:
                continue
            stats = parse_stats(fields.get("BaseStats")) or list(base.stats)
            types = parse_types(fields.get("Types")) or list(base.types)
            abilities = fields.get("Abilities", base.abilities.replace(" / ", ","))
            forms.append(
                SpeciesRow(
                    dex=base.dex,
                    species_id=sid,
                    form=form_n,
                    name=base.name,
                    form_name=fields.get("FormName") or f"Form {form_n}",
                    types=types,
                    stats=stats,
                    bst=sum(stats),
                    generation=fields.get("Generation", base.generation),
                    flags=fields.get("Flags", base.flags),
                    abilities=abilities.replace(",", " / "),
                    hidden_ability=fields.get("HiddenAbilities", base.hidden_ability),
                    sprite_key=sprite_stem(sid, form_n),
                )
            )
    return forms


def collect_species_rows() -> tuple[list[SpeciesRow], list[SpeciesRow], list[SpeciesRow]]:
    base = load_base_species()
    by_id = {r.species_id: r for r in base}
    forms = load_forms(by_id)
    forms_by_base: dict[str, list[SpeciesRow]] = defaultdict(list)
    for f in forms:
        forms_by_base[f.species_id].append(f)
    all_rows: list[SpeciesRow] = []
    for b in base:
        all_rows.append(b)
        for f in sorted(forms_by_base.get(b.species_id, []), key=lambda x: x.form):
            all_rows.append(f)
    return base, forms, all_rows


def build_catalog(rows: list[SpeciesRow]) -> dict[str, SpeciesRow]:
    """Lookup by catalog_key and by bare species id (base form)."""
    catalog: dict[str, SpeciesRow] = {}
    for row in rows:
        catalog[row.catalog_key] = row
        if row.form == 0:
            catalog[row.species_id] = row
    return catalog


def lookup_species(catalog: dict[str, SpeciesRow], raw_id: str) -> SpeciesRow | None:
    """Resolve encounter/trainer species tokens (BULBASAUR, SHELLOS_1, etc.)."""
    if raw_id in catalog:
        return catalog[raw_id]
    if "_" in raw_id:
        base, _, suffix = raw_id.partition("_")
        if suffix.isdigit():
            key = f"{base},{suffix}"
            if key in catalog:
                return catalog[key]
        if base in catalog:
            return catalog[base]
    return None


def type_fill(type_id: str) -> PatternFill:
    return PatternFill("solid", fgColor=TYPE_COLORS.get(type_id.upper(), "9CA3AF"))


def icon_png_bytes(stem: str, size: int = ICON_SIZE) -> bytes | None:
    path = resolve_sprite_path(stem, "Icons") or resolve_sprite_path(stem, "Front")
    if not path:
        return None
    im = PILImage.open(path).convert("RGBA")
    if path.parent.name == "Icons" and im.width >= im.height * 2:
        im = im.crop((0, 0, im.height, im.height))
    im = im.resize((size, size), PILImage.Resampling.NEAREST)
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def pokeapi_sprite_url(dex: int) -> str:
    return f"https://raw.githubusercontent.com/PokeAPI/sprites/master/sprites/pokemon/{dex}.png"


def sheets_sprite_formula(species_catalog_key: str) -> str:
    """Google Sheets formula that pulls Sprite from the Available Pokémon tab."""
    key = species_catalog_key.replace('"', '""')
    return (
        f'=IFERROR(INDEX(\'{POKEMON_TAB}\'!${POKEMON_SPRITE_COL}:${POKEMON_SPRITE_COL},'
        f'MATCH("{key}",\'{POKEMON_TAB}\'!${POKEMON_SPECIES_ID_COL}:${POKEMON_SPECIES_ID_COL},0)),"")'
    )


def fmt_name(raw: str) -> str:
    base, _, suffix = raw.partition("_")
    name = base.title()
    if suffix.isdigit():
        return f"{name} (Form {suffix})"
    if suffix:
        return f"{name} ({suffix.title()})"
    return name


def fmt_types(types: list[str]) -> str:
    return " / ".join(t.title() for t in types)


def fmt_moves(moves: list[str]) -> str:
    if not moves:
        return "— (level-up moveset)"
    return ", ".join(fmt_name(m) for m in moves)


def fmt_item(item: str | None) -> str:
    return f"@{fmt_name(item)}" if item else ""


def level_str(lv_min: int, lv_max: int | None) -> str:
    if lv_max is None or lv_max == lv_min:
        return str(lv_min)
    return f"{lv_min}–{lv_max}"


def apply_header_row(ws, row: int, headers: list[str], widths: list[int]):
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = CENTER
        cell.border = BORDER_ALL
        ws.column_dimensions[get_column_letter(col)].width = w


def style_data_cell(cell, alt: bool = False, bold: bool = False):
    cell.font = Font(name="Arial", size=10, bold=bold)
    cell.fill = FILL_ALT if alt else FILL_WHITE
    cell.border = BORDER_ALL
    cell.alignment = WRAP


def parse_trainer_types(path: Path) -> dict[str, str]:
    types: dict[str, str] = {}
    current_id: str | None = None
    for line in path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^\[([A-Z0-9_]+)\]$", line.strip())
        if m:
            current_id = m.group(1)
            continue
        if current_id and line.startswith("Name = "):
            types[current_id] = line.split("=", 1)[1].strip()
    return types


def parse_map_metadata(path: Path) -> dict[int, dict]:
    maps: dict[int, dict] = {}
    current_id: int | None = None
    current: dict = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^\[(\d+)\]", line.strip())
        if m:
            if current_id is not None:
                maps[current_id] = current
            current_id = int(m.group(1))
            current = {"id": current_id}
            cm = re.search(r"#\s*(.+)$", line)
            if cm:
                current["editor_name"] = cm.group(1).strip()
            continue
        if current_id is None or "=" not in line:
            continue
        key, val = [p.strip() for p in line.split("=", 1)]
        current[key] = val
    if current_id is not None:
        maps[current_id] = current
    return maps


def parse_encounters(path: Path, maps: dict[int, dict]) -> list[EncounterSlot]:
    slots: list[EncounterSlot] = []
    map_id: int | None = None
    version: int | None = None
    location = ""
    method = ""

    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.rstrip()
        if line.startswith("#") or not line.strip():
            continue
        hm = re.match(r"^\[(\d+)(?:,(\d+))?\]\s*#\s*(.+)$", line)
        if hm:
            map_id = int(hm.group(1))
            version = int(hm.group(2)) if hm.group(2) else None
            location = hm.group(3).strip()
            method = ""
            continue
        if map_id is None:
            continue
        # Method headers are not indented (Land,21 / Water,2 / OldRod).
        # Encounter slots are indented density lines (40,PIDGEY,11,14).
        if not line[:1].isspace():
            method = line.split(",", 1)[0].strip()
            continue
        parts = [p.strip() for p in line.split(",")]
        if len(parts) < 3:
            continue
        meta = maps.get(map_id, {})
        display = location or meta.get("Name", meta.get("editor_name", f"Map {map_id}"))
        if version is not None:
            display = f"{display} (v{version})"
        slots.append(
            EncounterSlot(
                map_id=map_id,
                version=version,
                location=display,
                method=method,
                rate=int(parts[0]),
                species=parts[1],
                level_min=int(parts[2]),
                level_max=int(parts[3]) if len(parts) > 3 else None,
            )
        )
    return slots


def parse_trainers(path: Path) -> list[Trainer]:
    trainers: list[Trainer] = []
    current: Trainer | None = None
    current_mon: TrainerMon | None = None

    def flush_mon():
        nonlocal current_mon
        if current and current_mon:
            current.pokemon.append(current_mon)
        current_mon = None

    def flush_trainer():
        nonlocal current
        flush_mon()
        if current:
            trainers.append(current)
        current = None

    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.rstrip()
        if line.startswith("#") or not line.strip():
            continue
        hm = re.match(r"^\[([^,\]]+),([^,\]]+)(?:,(\d+))?\]$", line.strip())
        if hm:
            flush_trainer()
            current = Trainer(
                trainer_type=hm.group(1),
                name=hm.group(2),
                version=int(hm.group(3)) if hm.group(3) else 0,
            )
            continue
        if current is None:
            continue
        if line.startswith("Pokemon = "):
            flush_mon()
            parts = [p.strip() for p in line.split("=", 1)[1].split(",")]
            current_mon = TrainerMon(species=parts[0], level=int(parts[1]))
            continue
        if current_mon is None:
            if line.startswith("LoseText = "):
                current.lose_text = line.split("=", 1)[1].strip()
            elif line.startswith("Items = "):
                current.items = [i.strip() for i in line.split("=", 1)[1].split(",")]
            continue
        key, val = [p.strip() for p in line.split("=", 1)]
        if key == "Name":
            current_mon.name = val
        elif key == "Moves":
            current_mon.moves = [m.strip() for m in val.split(",")]
        elif key == "Item":
            current_mon.item = val
        elif key == "AbilityIndex":
            current_mon.ability_index = val
        elif key == "Gender":
            current_mon.gender = val
        elif key == "IV":
            current_mon.ivs = val
        elif key == "Shiny":
            current_mon.shiny = val.lower() == "true"
        elif key == "Shadow":
            current_mon.shadow = val.lower() == "true"
        elif key == "Ball":
            current_mon.ball = val
    flush_trainer()
    return trainers


def _scan_for_battles(obj, map_id: int, event_id, trainer_placements, static_battles):
    pat_trainer = re.compile(
        r'TrainerBattle\.start\(\s*:(\w+)\s*,\s*"([^"]+)"(?:\s*,\s*(\d+))?\s*\)'
    )
    pat_wild = re.compile(r"WildBattle\.start\(:(\w+)\s*,\s*(\d+)\)")
    if isinstance(obj, str):
        for m in pat_trainer.finditer(obj):
            key = (m.group(1), m.group(2), int(m.group(3) or 0))
            trainer_placements[key].append({"map_id": map_id, "event_id": event_id})
        for m in pat_wild.finditer(obj):
            static_battles.append(
                {
                    "map_id": map_id,
                    "species": m.group(1),
                    "level": int(m.group(2)),
                    "event_id": event_id,
                }
            )
    elif isinstance(obj, list):
        for item in obj:
            _scan_for_battles(item, map_id, event_id, trainer_placements, static_battles)
    elif isinstance(obj, dict):
        for item in obj.values():
            _scan_for_battles(item, map_id, event_id, trainer_placements, static_battles)


def parse_battle_placements(maps_dir: Path) -> tuple[dict, list]:
    trainer_placements: dict[tuple, list[dict]] = defaultdict(list)
    static_battles: list[dict] = []
    if not maps_dir.is_dir():
        return trainer_placements, static_battles
    for path in sorted(maps_dir.glob("map*.json")):
        map_id = int(path.stem.replace("map", ""))
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            continue
        events = data.get("events") or {}
        event_iter = enumerate(events, 1) if isinstance(events, list) else events.items()
        for _key, ev in event_iter:
            if not ev or not isinstance(ev, dict):
                continue
            _scan_for_battles(ev, map_id, ev.get("id"), trainer_placements, static_battles)
    return trainer_placements, static_battles


def trainer_key(t: Trainer) -> tuple:
    return (t.trainer_type, t.name, t.version)


def map_sort_key(map_id: int, maps: dict[int, dict]) -> tuple:
    meta = maps.get(map_id, {})
    pos = meta.get("MapPosition", "0,99,99")
    parts = [int(p) for p in pos.split(",")]
    region = parts[0] if parts else 0
    x = parts[1] if len(parts) > 1 else 99
    y = parts[2] if len(parts) > 2 else 99
    return (region, y, x, map_id, meta.get("Name", ""))


def trainer_display_name(t: Trainer, trainer_types: dict[str, str]) -> str:
    cls = trainer_types.get(t.trainer_type, fmt_name(t.trainer_type))
    suffix = f" #{t.version + 1}" if t.version else ""
    return f"{cls} {t.name}{suffix}"


def mon_detail_line(mon: TrainerMon, catalog: dict[str, SpeciesRow]) -> str:
    sp = lookup_species(catalog, mon.species)
    name = mon.name or (sp.display_name if sp else fmt_name(mon.species))
    bst = f" BST {sp.bst}" if sp else ""
    types = f" [{fmt_types(sp.types)}]" if sp else ""
    parts = [f"{name} Lv.{mon.level}{types}{bst}"]
    if mon.item:
        parts.append(fmt_item(mon.item))
    parts.append(fmt_moves(mon.moves))
    extras = []
    if mon.gender:
        extras.append(mon.gender.title())
    if mon.shiny:
        extras.append("Shiny")
    if mon.shadow:
        extras.append("Shadow")
    if mon.ivs:
        extras.append(f"IV {mon.ivs}")
    if mon.ability_index is not None:
        extras.append(f"Ability slot {mon.ability_index}")
    if mon.ball:
        extras.append(fmt_name(mon.ball))
    if extras:
        parts.append(f"[{' | '.join(extras)}]")
    return " ".join(parts)
