#!/usr/bin/env python3
"""
Generate Run & Bun / Kaizo-style battle documentation spreadsheets from PBS data.

Output: docs/sheets/Essentials_Remix_AI_Battles.xlsx
Import into Google Sheets: File → Import → Upload → Replace spreadsheet or create new.
"""

from __future__ import annotations

import json
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
PBS = ROOT / "PBS"
OUT = ROOT / "docs" / "sheets" / "Essentials_Remix_AI_Battles.xlsx"

# Pokémon type colors (hex without #) — standard community palette
TYPE_COLORS = {
    "NORMAL": "A8A878",
    "FIRE": "F08030",
    "WATER": "6890F0",
    "ELECTRIC": "F8D030",
    "GRASS": "78C850",
    "ICE": "98D8D8",
    "FIGHTING": "C03028",
    "POISON": "A040A0",
    "GROUND": "E0C068",
    "FLYING": "A890F0",
    "PSYCHIC": "F85888",
    "BUG": "A8B820",
    "ROCK": "B8A038",
    "GHOST": "705898",
    "DRAGON": "7038F8",
    "DARK": "705848",
    "STEEL": "B8B8D0",
    "FAIRY": "EE99AC",
}

METHOD_LABELS = {
    "Land": "Grass",
    "LandNight": "Grass (Night)",
    "LandMorning": "Grass (Morning)",
    "LandDay": "Grass (Day)",
    "Water": "Surf",
    "OldRod": "Old Rod",
    "GoodRod": "Good Rod",
    "SuperRod": "Super Rod",
    "Cave": "Cave",
    "RockSmash": "Rock Smash",
    "HeadbuttLow": "Headbutt (Low)",
    "HeadbuttHigh": "Headbutt (High)",
    "PokeRadar": "Poké Radar",
    "BugContest": "Bug Contest",
}

METHOD_ORDER = list(METHOD_LABELS.keys()) + ["Other"]

# Styles
FONT_TITLE = Font(name="Arial", size=18, bold=True, color="FFFFFF")
FONT_SECTION = Font(name="Arial", size=12, bold=True, color="FFFFFF")
FONT_HEADER = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FONT_BODY = Font(name="Arial", size=10)
FONT_TRAINER = Font(name="Arial", size=11, bold=True, color="1F2937")
FONT_NOTE = Font(name="Arial", size=9, italic=True, color="6B7280")

FILL_TITLE = PatternFill("solid", fgColor="1E3A5F")
FILL_SECTION = PatternFill("solid", fgColor="2563EB")
FILL_HEADER = PatternFill("solid", fgColor="374151")
FILL_ALT = PatternFill("solid", fgColor="F3F4F6")
FILL_WHITE = PatternFill("solid", fgColor="FFFFFF")
FILL_BOSS = PatternFill("solid", fgColor="FEF3C7")

THIN = Side(style="thin", color="D1D5DB")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def fmt_name(raw: str) -> str:
    """BULBASAUR / SHELLOS_1 → display name."""
    base, _, suffix = raw.partition("_")
    name = base.title()
    if suffix.isdigit():
        return f"{name} (Form {suffix})"
    if suffix:
        return f"{name} ({suffix.title()})"
    return name


def fmt_types(types: list[str]) -> str:
    return " / ".join(t.title() for t in types)


def fmt_moves(moves: list[str]) -> str:
    if not moves:
        return "— (level-up moveset)"
    return ", ".join(fmt_name(m) for m in moves)


def fmt_item(item: str | None) -> str:
    if not item:
        return ""
    return f"@{fmt_name(item)}"


def parse_pokemon_db(path: Path) -> dict[str, dict]:
    species: dict[str, dict] = {}
    current: dict | None = None
    for line in path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^\[([A-Z0-9_]+)\]$", line.strip())
        if m:
            current = {"id": m.group(1), "types": []}
            species[m.group(1)] = current
            continue
        if current is None or "=" not in line:
            continue
        key, val = [p.strip() for p in line.split("=", 1)]
        if key == "Name":
            current["name"] = val
        elif key == "Types":
            current["types"] = [t.strip() for t in val.split(",")]
    return species


def parse_trainer_types(path: Path) -> dict[str, str]:
    types: dict[str, str] = {}
    current_id: str | None = None
    for line in path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^\[([A-Z0-9_]+)\]$", line.strip())
        if m:
            current_id = m.group(1)
            continue
        if current_id and line.startswith("Name = "):
            types[current_id] = line.split("=", 1)[1].strip()
    return types


def parse_map_metadata(path: Path) -> dict[int, dict]:
    maps: dict[int, dict] = {}
    current_id: int | None = None
    current: dict = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^\[(\d+)\]", line.strip())
        if m:
            if current_id is not None:
                maps[current_id] = current
            current_id = int(m.group(1))
            current = {"id": current_id}
            cm = re.search(r"#\s*(.+)$", line)
            if cm:
                current["editor_name"] = cm.group(1).strip()
            continue
        if current_id is None or "=" not in line:
            continue
        key, val = [p.strip() for p in line.split("=", 1)]
        current[key] = val
    if current_id is not None:
        maps[current_id] = current
    return maps


@dataclass
class EncounterSlot:
    map_id: int
    version: int | None
    location: str
    method: str
    rate: int
    species: str
    level_min: int
    level_max: int | None


def parse_encounters(path: Path, maps: dict[int, dict]) -> list[EncounterSlot]:
    slots: list[EncounterSlot] = []
    map_id: int | None = None
    version: int | None = None
    location = ""
    method = ""
    step_rate: int | None = None

    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.rstrip()
        if line.startswith("#") or not line.strip():
            continue
        hm = re.match(r"^\[(\d+)(?:,(\d+))?\]\s*#\s*(.+)$", line)
        if hm:
            map_id = int(hm.group(1))
            version = int(hm.group(2)) if hm.group(2) else None
            location = hm.group(3).strip()
            method = ""
            step_rate = None
            continue
        if map_id is None:
            continue
        if "," not in line and not line.startswith(" "):
            parts = line.split(",", 1)
            method = parts[0].strip()
            step_rate = int(parts[1]) if len(parts) > 1 else None
            continue
        parts = [p.strip() for p in line.split(",")]
        if len(parts) < 3:
            continue
        rate = int(parts[0])
        species = parts[1]
        lv_min = int(parts[2])
        lv_max = int(parts[3]) if len(parts) > 3 else None
        meta = maps.get(map_id, {})
        display = location or meta.get("Name", meta.get("editor_name", f"Map {map_id}"))
        if version is not None:
            display = f"{display} (v{version})"
        slots.append(
            EncounterSlot(
                map_id=map_id,
                version=version,
                location=display,
                method=method,
                rate=rate,
                species=species,
                level_min=lv_min,
                level_max=lv_max,
            )
        )
    return slots


@dataclass
class TrainerMon:
    species: str
    level: int
    name: str | None = None
    moves: list[str] = field(default_factory=list)
    item: str | None = None
    ability_index: str | None = None
    gender: str | None = None
    ivs: str | None = None
    shiny: bool = False
    shadow: bool = False
    ball: str | None = None


@dataclass
class Trainer:
    trainer_type: str
    name: str
    version: int
    lose_text: str = ""
    items: list[str] = field(default_factory=list)
    pokemon: list[TrainerMon] = field(default_factory=list)


def parse_trainers(path: Path) -> list[Trainer]:
    trainers: list[Trainer] = []
    current: Trainer | None = None
    current_mon: TrainerMon | None = None

    def flush_mon():
        nonlocal current_mon
        if current and current_mon:
            current.pokemon.append(current_mon)
        current_mon = None

    def flush_trainer():
        nonlocal current
        flush_mon()
        if current:
            trainers.append(current)
        current = None

    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.rstrip()
        if line.startswith("#") or not line.strip():
            continue
        hm = re.match(r"^\[([^,\]]+),([^,\]]+)(?:,(\d+))?\]$", line.strip())
        if hm:
            flush_trainer()
            current = Trainer(
                trainer_type=hm.group(1),
                name=hm.group(2),
                version=int(hm.group(3)) if hm.group(3) else 0,
            )
            continue
        if current is None:
            continue
        if line.startswith("Pokemon = "):
            flush_mon()
            parts = [p.strip() for p in line.split("=", 1)[1].split(",")]
            current_mon = TrainerMon(species=parts[0], level=int(parts[1]))
            continue
        if current_mon is None:
            if line.startswith("LoseText = "):
                current.lose_text = line.split("=", 1)[1].strip()
            elif line.startswith("Items = "):
                current.items = [i.strip() for i in line.split("=", 1)[1].split(",")]
            continue
        key, val = [p.strip() for p in line.split("=", 1)]
        if key == "Name":
            current_mon.name = val
        elif key == "Moves":
            current_mon.moves = [m.strip() for m in val.split(",")]
        elif key == "Item":
            current_mon.item = val
        elif key == "AbilityIndex":
            current_mon.ability_index = val
        elif key == "Gender":
            current_mon.gender = val
        elif key == "IV":
            current_mon.ivs = val
        elif key == "Shiny":
            current_mon.shiny = val.lower() == "true"
        elif key == "Shadow":
            current_mon.shadow = val.lower() == "true"
        elif key == "Ball":
            current_mon.ball = val
    flush_trainer()
    return trainers


def _scan_for_battles(obj, map_id: int, event_id, trainer_placements, static_battles):
    """Recursively find battle script strings in map event JSON."""
    pat_trainer = re.compile(
        r'TrainerBattle\.start\(\s*:(\w+)\s*,\s*"([^"]+)"(?:\s*,\s*(\d+))?\s*\)'
    )
    pat_wild = re.compile(r"WildBattle\.start\(:(\w+)\s*,\s*(\d+)\)")

    if isinstance(obj, str):
        for m in pat_trainer.finditer(obj):
            key = (m.group(1), m.group(2), int(m.group(3) or 0))
            trainer_placements[key].append({"map_id": map_id, "event_id": event_id})
        for m in pat_wild.finditer(obj):
            static_battles.append(
                {
                    "map_id": map_id,
                    "species": m.group(1),
                    "level": int(m.group(2)),
                    "event_id": event_id,
                }
            )
    elif isinstance(obj, list):
        for item in obj:
            _scan_for_battles(item, map_id, event_id, trainer_placements, static_battles)
    elif isinstance(obj, dict):
        for item in obj.values():
            _scan_for_battles(item, map_id, event_id, trainer_placements, static_battles)


def parse_battle_placements(maps_dir: Path) -> tuple[dict, list]:
    """Map TrainerBattle.start / WildBattle.start to map locations."""
    trainer_placements: dict[tuple, list[dict]] = defaultdict(list)
    static_battles: list[dict] = []

    for path in sorted(maps_dir.glob("map*.json")):
        map_id = int(path.stem.replace("map", ""))
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            continue
        events = data.get("events") or {}
        if isinstance(events, list):
            event_iter = enumerate(events, 1)
        else:
            event_iter = events.items()
        for _key, ev in event_iter:
            if not ev or not isinstance(ev, dict):
                continue
            event_id = ev.get("id")
            _scan_for_battles(ev, map_id, event_id, trainer_placements, static_battles)

    return trainer_placements, static_battles


def trainer_key(t: Trainer) -> tuple:
    return (t.trainer_type, t.name, t.version)


def map_sort_key(map_id: int, maps: dict[int, dict]) -> tuple:
    """Order locations roughly by region map position (north→south, west→east)."""
    meta = maps.get(map_id, {})
    pos = meta.get("MapPosition", "0,99,99")
    parts = [int(p) for p in pos.split(",")]
    region = parts[0] if parts else 0
    x = parts[1] if len(parts) > 1 else 99
    y = parts[2] if len(parts) > 2 else 99
    name = meta.get("Name", "")
    return (region, y, x, map_id, name)


def level_str(lv_min: int, lv_max: int | None) -> str:
    if lv_max is None or lv_max == lv_min:
        return str(lv_min)
    return f"{lv_min}–{lv_max}"


def apply_header_row(ws, row: int, headers: list[str], widths: list[int]):
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = CENTER
        cell.border = BORDER_ALL
        ws.column_dimensions[get_column_letter(col)].width = w


def style_data_cell(cell, alt: bool = False, bold: bool = False):
    cell.font = Font(name="Arial", size=10, bold=bold)
    cell.fill = FILL_ALT if alt else FILL_WHITE
    cell.border = BORDER_ALL
    cell.alignment = WRAP


def type_fill(primary_type: str) -> PatternFill:
    color = TYPE_COLORS.get(primary_type.upper(), "E5E7EB")
    return PatternFill("solid", fgColor=color)


def write_index_sheet(wb: Workbook, encounter_count: int, trainer_count: int):
    ws = wb.active
    ws.title = "Index"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = "Essentials Remix AI — Battle Documentation"
    c.font = FONT_TITLE
    c.fill = FILL_TITLE
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    rows = [
        ("", ""),
        ("Game version", "1.0.0 (Pokémon Essentials 21.1, Gen 8 mechanics)"),
        ("Generated", str(date.today())),
        ("Source", "Auto-exported from PBS/ (encounters.txt, trainers.txt, map events)"),
        ("", ""),
        ("Sheets", ""),
        ("Wild Encounters", f"{encounter_count} encounter slots across all routes and areas"),
        ("Trainer Battles", f"{trainer_count} trainer definitions (PBS + map placements)"),
        ("Encounter Summary", "One row per area — quick species checklist"),
        ("Static Battles", "Scripted wild battles (legendaries, etc.)"),
        ("", ""),
        ("How to use", ""),
        (
            "Encounters",
            "Grouped by location. Rate % is the slot weight within that method "
            "(higher = more common). Method shows grass, surf, rods, etc.",
        ),
        (
            "Trainers",
            "Grouped by map location when known. Moves shown when defined in PBS; "
            "otherwise the game uses the species' level-up moveset.",
        ),
        (
            "Regenerate",
            "Run: python3 tools/generate_battle_docs.py — then re-import the .xlsx into Google Sheets.",
        ),
        ("", ""),
        ("Legend", ""),
        ("— (dash)", "Data not specified in PBS (uses game defaults)"),
        ("[Boss]", "Gym Leader, Champion, or flagged boss fights (fill in as you add them)"),
        ("(vN)", "Alternate encounter table version for the same map"),
    ]
    for i, (k, v) in enumerate(rows, 3):
        ws.cell(i, 1, k).font = Font(name="Arial", size=10, bold=bool(k and k not in ("", "Sheets", "How to use", "Legend")))
        ws.cell(i, 2, v).font = FONT_BODY
        ws.cell(i, 2).alignment = WRAP
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 72


def write_encounters_sheet(wb: Workbook, slots: list[EncounterSlot], species_db: dict, maps: dict):
    ws = wb.create_sheet("Wild Encounters")
    headers = [
        "Location",
        "Map",
        "Method",
        "Pokémon",
        "Type",
        "Level",
        "Rate %",
        "Notes",
    ]
    widths = [28, 8, 16, 18, 14, 10, 10, 24]
    apply_header_row(ws, 1, headers, widths)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H1"

    # Group by map_id + version + location
    groups: dict[tuple, list[EncounterSlot]] = defaultdict(list)
    for s in slots:
        groups[(s.map_id, s.version or 0, s.location)].append(s)

    row = 2
    sorted_keys = sorted(
        groups.keys(),
        key=lambda k: (map_sort_key(k[0], maps), k[1], k[2]),
    )
    for key in sorted_keys:
        map_id, _ver, location = key
        group_slots = groups[key]

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        sec = ws.cell(row, 1, f"▸  {location}  (Map {map_id})")
        sec.font = FONT_SECTION
        sec.fill = FILL_SECTION
        sec.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 22
        row += 1

        by_method: dict[str, list[EncounterSlot]] = defaultdict(list)
        for s in group_slots:
            by_method[s.method].append(s)

        alt = False
        sorted_methods = sorted(
            by_method.keys(),
            key=lambda m: (METHOD_ORDER.index(m) if m in METHOD_ORDER else 99, m),
        )
        for method in sorted_methods:
            method_label = METHOD_LABELS.get(method, method)
            for s in sorted(by_method[method], key=lambda x: (-x.rate, x.species)):
                sp = species_db.get(s.species, {})
                types = sp.get("types", [])
                type_str = fmt_types(types) if types else "?"
                display_name = sp.get("name") or fmt_name(s.species)

                values = [
                    location,
                    map_id,
                    method_label,
                    display_name,
                    type_str,
                    level_str(s.level_min, s.level_max),
                    s.rate,
                    "",
                ]
                for col, val in enumerate(values, 1):
                    cell = ws.cell(row, col, val)
                    style_data_cell(cell, alt=alt)
                    if col == 5 and types:
                        cell.fill = type_fill(types[0])
                        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                alt = not alt
                row += 1

    return len(sorted_keys)


def write_location_summary(wb: Workbook, slots: list[EncounterSlot], species_db: dict, maps: dict):
    """Quick-reference: one row per location with species list."""
    ws = wb.create_sheet("Encounter Summary")
    headers = ["Location", "Map", "Methods", "Species (unique)", "Level Range", "Slot Count"]
    widths = [28, 8, 24, 40, 14, 12]
    apply_header_row(ws, 1, headers, widths)
    ws.freeze_panes = "A2"

    by_loc: dict[tuple, list[EncounterSlot]] = defaultdict(list)
    for s in slots:
        by_loc[(s.map_id, s.location)].append(s)

    row = 2
    alt = False
    for (map_id, location) in sorted(by_loc.keys(), key=lambda k: map_sort_key(k[0], maps)):
        group = by_loc[(map_id, location)]
        methods = sorted({METHOD_LABELS.get(s.method, s.method) for s in group})
        species = sorted(
            {species_db.get(s.species, {}).get("name") or fmt_name(s.species) for s in group}
        )
        levels = []
        for s in group:
            levels.append(s.level_min)
            if s.level_max:
                levels.append(s.level_max)
        lv_range = f"{min(levels)}–{max(levels)}" if levels else "—"

        values = [
            location,
            map_id,
            ", ".join(methods),
            ", ".join(species),
            lv_range,
            len(group),
        ]
        for col, val in enumerate(values, 1):
            style_data_cell(ws.cell(row, col, val), alt=alt)
        alt = not alt
        row += 1


def trainer_display_name(t: Trainer, trainer_types: dict[str, str]) -> str:
    cls = trainer_types.get(t.trainer_type, fmt_name(t.trainer_type))
    suffix = f" #{t.version + 1}" if t.version else ""
    if t.trainer_type.startswith("LEADER_") or t.trainer_type == "CHAMPION":
        return f"{cls} {t.name}{suffix}"
    return f"{cls} {t.name}{suffix}"


def mon_detail_line(mon: TrainerMon, species_db: dict) -> str:
    sp = species_db.get(mon.species, {})
    name = mon.name or sp.get("name") or fmt_name(mon.species)
    parts = [f"{name} Lv.{mon.level}"]
    if mon.item:
        parts.append(fmt_item(mon.item))
    if mon.moves:
        parts.append(fmt_moves(mon.moves))
    else:
        parts.append("— (level-up moveset)")
    extras = []
    if mon.gender:
        extras.append(mon.gender.title())
    if mon.shiny:
        extras.append("Shiny")
    if mon.shadow:
        extras.append("Shadow")
    if mon.ivs:
        extras.append(f"IV {mon.ivs}")
    if mon.ability_index is not None:
        extras.append(f"Ability slot {mon.ability_index}")
    if mon.ball:
        extras.append(fmt_name(mon.ball))
    if extras:
        parts.append(f"[{' | '.join(extras)}]")
    return " ".join(parts)


def write_trainers_sheet(
    wb: Workbook,
    trainers: list[Trainer],
    trainer_types: dict[str, str],
    species_db: dict,
    placements: dict,
    maps: dict[int, dict],
):
    ws = wb.create_sheet("Trainer Battles")
    headers = [
        "Location",
        "Map",
        "Trainer",
        "Party",
        "Trainer Items",
        "Lose Text",
        "Notes",
    ]
    widths = [26, 8, 24, 56, 22, 28, 20]
    apply_header_row(ws, 1, headers, widths)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G1"

    # Build placement lookup per trainer
    placed: dict[tuple, list[dict]] = defaultdict(list)
    unplaced: list[Trainer] = []
    for t in trainers:
        key = trainer_key(t)
        locs = placements.get(key, [])
        if locs:
            placed[key].extend(locs)
        else:
            unplaced.append(t)

    # Ordered: by map id, then trainer
    entries: list[tuple] = []
    seen_keys: set[tuple] = set()
    for t in trainers:
        key = trainer_key(t)
        if key in seen_keys:
            continue
        seen_keys.add(key)
        locs = placed.get(key, [])
        if locs:
            for loc in locs:
                entries.append((loc["map_id"], t, loc))
        else:
            entries.append((9999, t, None))

    entries.sort(
        key=lambda e: (
            map_sort_key(e[0], maps) if e[0] != 9999 else (9, 99, 99, 9999, ""),
            e[1].trainer_type,
            e[1].name,
            e[1].version,
        )
    )

    row = 2
    current_map: int | None = None
    alt = False
    for map_id, trainer, loc in entries:
        if map_id != current_map:
            current_map = map_id
            if map_id != 9999:
                meta = maps.get(map_id, {})
                loc_name = meta.get("Name", meta.get("editor_name", f"Map {map_id}"))
            else:
                loc_name = "Unplaced / Script-Only Battles"
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
            sec = ws.cell(row, 1, f"▸  {loc_name}")
            sec.font = FONT_SECTION
            sec.fill = FILL_SECTION
            ws.row_dimensions[row].height = 22
            row += 1
            alt = False

        is_boss = trainer.trainer_type.startswith("LEADER_") or trainer.trainer_type == "CHAMPION"
        meta = maps.get(map_id, {}) if map_id != 9999 else {}
        loc_name = meta.get("Name", "") if map_id != 9999 else "—"
        party_lines = [mon_detail_line(m, species_db) for m in trainer.pokemon]
        party = "\n".join(party_lines)
        t_items = ", ".join(fmt_name(i) for i in trainer.items) if trainer.items else "—"
        notes = ""
        if is_boss:
            notes = "Boss"
        if loc and loc.get("event_id"):
            notes = (notes + f" Event {loc['event_id']}").strip()

        values = [
            loc_name,
            map_id if map_id != 9999 else "—",
            trainer_display_name(trainer, trainer_types),
            party,
            t_items,
            trainer.lose_text or "—",
            notes,
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row, col, val)
            fill = FILL_BOSS if is_boss else (FILL_ALT if alt else FILL_WHITE)
            cell.fill = fill
            cell.font = FONT_TRAINER if col == 3 else FONT_BODY
            cell.border = BORDER_ALL
            cell.alignment = WRAP
        alt = not alt
        row += 1

    # Script-only battles documented in Index / notes
    script_notes = [
        ("Northshire Abbey", "Kobold Worker / Laborer / Vermin Scout", "Quest Journal — dynamic parties"),
        ("Northshire Abbey", "Defias Thugs", "Quest chain battles"),
    ]
    if script_notes:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        sec = ws.cell(row, 1, "▸  Script-Generated Battles (not in trainers.txt)")
        sec.font = FONT_SECTION
        sec.fill = PatternFill("solid", fgColor="7C3AED")
        row += 1
        for loc, trainer, note in script_notes:
            for col, val in enumerate([loc, "—", trainer, note, "—", "—", "Plugin"], 1):
                cell = ws.cell(row, col, val)
                style_data_cell(cell, alt=alt)
                cell.font = FONT_NOTE if col == 4 else FONT_BODY
            alt = not alt
            row += 1


def write_static_sheet(wb: Workbook, static: list, species_db: dict, maps: dict):
    ws = wb.create_sheet("Static Battles")
    headers = ["Location", "Map", "Pokémon", "Type", "Level", "Notes"]
    widths = [28, 8, 18, 14, 10, 30]
    apply_header_row(ws, 1, headers, widths)
    ws.freeze_panes = "A2"

    row = 2
    alt = False
    for b in sorted(static, key=lambda x: x["map_id"]):
        map_id = b["map_id"]
        meta = maps.get(map_id, {})
        loc = meta.get("Name", meta.get("editor_name", f"Map {map_id}"))
        sp = species_db.get(b["species"], {})
        types = sp.get("types", [])
        values = [
            loc,
            map_id,
            sp.get("name") or fmt_name(b["species"]),
            fmt_types(types) if types else "?",
            b["level"],
            f"Event {b['event_id']} — WildBattle.start",
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row, col, val)
            style_data_cell(cell, alt=alt)
            if col == 4 and types:
                cell.fill = type_fill(types[0])
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        alt = not alt
        row += 1


def main():
    species_db = parse_pokemon_db(PBS / "pokemon.txt")
    trainer_types = parse_trainer_types(PBS / "trainer_types.txt")
    maps = parse_map_metadata(PBS / "map_metadata.txt")
    encounters = parse_encounters(PBS / "encounters.txt", maps)
    trainers = parse_trainers(PBS / "trainers.txt")
    placements, static = parse_battle_placements(PBS / "MapEvents")

    OUT.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    write_index_sheet(wb, len(encounters), len(trainers))
    write_encounters_sheet(wb, encounters, species_db, maps)
    write_location_summary(wb, encounters, species_db, maps)
    write_trainers_sheet(wb, trainers, trainer_types, species_db, placements, maps)
    write_static_sheet(wb, static, species_db, maps)

    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"  Wild encounter slots: {len(encounters)}")
    print(f"  Trainer definitions:    {len(trainers)}")
    print(f"  Map-placed trainers:    {sum(len(v) for v in placements.values())}")
    print(f"  Static wild battles:    {len(static)}")


if __name__ == "__main__":
    main()
